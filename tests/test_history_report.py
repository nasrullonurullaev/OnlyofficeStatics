from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook, Workbook
from history_report import update_report, load_history, HistoryError, PRODUCTS
from test_desktop_downloads import report, RELEASE, asset


def product_data(release_id=10, tag='v10.0.0', count=100, identifier=1, name='DesktopEditors_x64.exe'):
    release = dict(RELEASE, id=release_id, tag_name=tag,
                   published_at=f'2026-10-{release_id:02d}T00:00:00Z')
    item = asset(name, count, identifier)
    item['category'] = report.classify(name)
    return release, [item]


class HistoryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'ONLYOFFICE-Download-Statistics.xlsx'
        self.now = datetime(2026, 10, 28, 9, 17, tzinfo=timezone.utc)

    def test_new_release_and_artifact_shift_rows_without_moving_old_values_to_other_files(self):
        old = product_data()
        products = {'Desktop Editors': [old], 'Document Server': []}
        update_report(self.path, products, self.now)
        newer = product_data(11, 'v11.0.0', 50, 3)
        existing = product_data(count=150)
        added = asset('DesktopEditors_arm64.exe', 7, 2)
        added['category'] = report.classify(added['name'])
        existing[1].insert(0, added)
        products['Desktop Editors'] = [existing, newer]
        update_report(self.path, products, self.now, history_path=self.path)
        state = load_history(self.path)
        self.assertEqual(state['assets']['Desktop Editors', '10', '1']['values'], [100, 150])
        self.assertEqual(state['assets']['Desktop Editors', '10', '2']['values'], [None, 7])
        self.assertEqual(state['assets']['Desktop Editors', '11', '3']['values'], [None, 50])
        book = load_workbook(self.path, data_only=True)
        self.assertEqual([s.title for s in book if s.sheet_state == 'visible'], list(PRODUCTS))
        sheet = book['Desktop Editors']
        rows = {cell.value: cell.row for row in sheet for cell in row if cell.column == 1}
        self.assertLess(rows['v11.0.0'], rows['v10.0.0'])
        new_row = rows['DesktopEditors_arm64.exe']
        self.assertEqual([sheet.cell(new_row, column).value for column in (2, 3)], ['—', 7])
        version_row = rows['v10.0.0']
        old_row = next(row for row in range(version_row, sheet.max_row + 1)
                       if sheet.cell(row, 1).value == 'DesktopEditors_x64.exe')
        self.assertEqual([sheet.cell(old_row, column).value for column in (2, 3)], [100, 150])
        book.close()

    def test_all_releases_remain_in_one_sheet_not_just_latest_two(self):
        releases = [product_data(i, f'v{i}.0.0', i * 10, i) for i in (7, 9, 8, 10)]
        update_report(self.path, {'Desktop Editors': releases, 'Document Server': []}, self.now)
        book = load_workbook(self.path, data_only=True)
        versions = [cell.value for row in book['Desktop Editors'] for cell in row
                    if cell.column == 1 and isinstance(cell.value, str) and cell.value.startswith('v')]
        self.assertEqual(versions, ['v10.0.0', 'v9.0.0', 'v8.0.0', 'v7.0.0'])
        book.close()

    def test_same_name_reuploaded_with_new_id_does_not_inherit_old_counts(self):
        products = {'Desktop Editors': [product_data(count=100, identifier=1)]}
        update_report(self.path, products, self.now)
        products = {'Desktop Editors': [product_data(count=5, identifier=2)]}
        update_report(self.path, products, self.now, history_path=self.path)
        state = load_history(self.path)
        self.assertEqual(state['assets']['Desktop Editors', '10', '1']['values'], [100, None])
        self.assertEqual(state['assets']['Desktop Editors', '10', '2']['values'], [None, 5])

    def test_missing_asset_is_preserved_and_zero_is_not_missing(self):
        products = {'Desktop Editors': [product_data(count=0)]}
        update_report(self.path, products, self.now)
        release = products['Desktop Editors'][0][0]
        products['Desktop Editors'] = [(release, [])]
        update_report(self.path, products, self.now, history_path=self.path)
        state = load_history(self.path)
        self.assertEqual(state['assets']['Desktop Editors', '10', '1']['values'], [0, None])
        book = load_workbook(self.path, data_only=True)
        totals = [row for row in book['Desktop Editors'] if row[0].value == 'Total Windows']
        self.assertEqual([cell.value for cell in totals[0][1:]], [0, '—'])
        book.close()

    def test_repeated_date_appends_columns_not_tabs(self):
        products = {'Desktop Editors': [product_data()]}
        for index in range(3):
            update_report(self.path, products, self.now, history_path=self.path if index else None)
        book = load_workbook(self.path)
        self.assertEqual(book.sheetnames, ['Desktop Editors', 'Document Server', '_history'])
        headers = [cell.value for cell in book['Desktop Editors'][4]]
        self.assertEqual(len(headers), 4)
        self.assertEqual(len(set(headers)), 4)
        self.assertEqual(book['_history'].sheet_state, 'hidden')
        book.close()

    def test_corrupt_or_old_history_is_not_overwritten(self):
        products = {'Desktop Editors': [product_data()]}
        self.path.write_bytes(b'bad file')
        with self.assertRaises(HistoryError):
            update_report(self.path, products, self.now, history_path=self.path)
        self.assertEqual(self.path.read_bytes(), b'bad file')
        book = Workbook()
        book.save(self.path)
        old = self.path.read_bytes()
        with self.assertRaisesRegex(HistoryError, 'old format'):
            update_report(self.path, products, self.now, history_path=self.path)
        self.assertEqual(self.path.read_bytes(), old)

    def test_products_are_independent_even_with_equal_release_and_asset_ids(self):
        products = {'Desktop Editors': [product_data(count=10)],
                    'Document Server': [product_data(count=90, name='onlyoffice-documentserver.exe')]}
        update_report(self.path, products, self.now)
        state = load_history(self.path)
        self.assertEqual(state['assets']['Desktop Editors', '10', '1']['values'], [10])
        self.assertEqual(state['assets']['Document Server', '10', '1']['values'], [90])


class CollectionTests(unittest.TestCase):
    def test_publishing_downloads_existing_matrix_and_adds_a_column(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old = root / 'old.xlsx'
            update_report(old, {'Desktop Editors': [product_data(count=10)]},
                          datetime(2026, 10, 21, tzinfo=timezone.utc))
            portal = Mock()
            portal.find.return_value = {'id': 5, 'version': 1}
            portal.download.side_effect = lambda target, info: target.write_bytes(old.read_bytes())
            output = root / 'ONLYOFFICE-Download-Statistics.xlsx'
            with patch('sys.argv', ['report', '--publish', '--output', str(output)]), \
                    patch.dict('os.environ', {'DOCSPACE_API_KEY': 'test'}), \
                    patch.object(report, 'DocSpace', return_value=portal), \
                    patch.object(report, 'collect_tracked', side_effect=[
                        [product_data(count=20)], [product_data(count=30, name='onlyoffice-documentserver.exe')]]):
                self.assertEqual(report.main(), 0)
            state = load_history(output)
            self.assertEqual(state['assets']['Desktop Editors', '10', '1']['values'], [10, 20])
            self.assertEqual(len(state['runs']), 2)
            portal.publish.assert_called_once_with(output, portal.find.return_value)

    def test_selects_stable_releases_with_pagination_and_skips_drafts(self):
        first = [dict(RELEASE, id=i, tag_name=f'v{i}', draft=True) for i in range(100)]
        older = dict(RELEASE, id=9, tag_name='v9', published_at='2026-09-01T00:00:00Z')
        beta = dict(RELEASE, id=11, prerelease=True)
        with patch.object(report, 'request_json', side_effect=[first, [older, beta, RELEASE]]) as request, \
                patch.object(report, 'collect_assets', return_value=[]) as collect:
            releases = report.collect_tracked(Mock())
        self.assertEqual([release['id'] for release, assets in releases], [10, 9])
        self.assertEqual(request.call_args.kwargs['params']['page'], 2)
        self.assertEqual(collect.call_count, 2)

    def test_first_run_selects_only_latest_two(self):
        releases = [product_data(i)[0] for i in (7, 10, 8, 9)]
        with patch.object(report, 'request_json', return_value=releases), \
                patch.object(report, 'collect_assets', return_value=[]):
            result = report.collect_tracked(Mock())
        self.assertEqual([release['id'] for release, _ in result], [10, 9])

    def test_later_runs_refresh_tracked_and_add_all_new_releases(self):
        releases = [product_data(i)[0] for i in (7, 9, 10, 11, 12)]
        known = {'9': {'published': releases[1]['published_at']},
                 '10': {'published': releases[2]['published_at']}}
        with patch.object(report, 'request_json', return_value=releases), \
                patch.object(report, 'collect_assets', return_value=[]):
            result = report.collect_tracked(Mock(), known_releases=known)
        self.assertEqual([release['id'] for release, _ in result], [12, 11, 10, 9])

    def test_no_new_release_does_not_import_older_versions(self):
        releases = [product_data(i)[0] for i in (7, 9, 10)]
        known = {str(item['id']): {'published': item['published_at']} for item in releases[1:]}
        with patch.object(report, 'request_json', return_value=releases), \
                patch.object(report, 'collect_assets', return_value=[]):
            result = report.collect_tracked(Mock(), known_releases=known)
        self.assertEqual([release['id'] for release, _ in result], [10, 9])

    def test_failed_collection_does_not_publish(self):
        portal = Mock()
        portal.find.return_value = None
        with patch('sys.argv', ['report', '--publish']), \
                patch.dict('os.environ', {'DOCSPACE_API_KEY': 'test'}), \
                patch.object(report, 'DocSpace', return_value=portal), \
                patch.object(report, 'collect_tracked', side_effect=report.ReportError('Incomplete collection')), \
                self.assertLogs(report.LOG, level='ERROR'):
            self.assertEqual(report.main(), 1)
        portal.publish.assert_not_called()


if __name__ == '__main__':
    unittest.main()
