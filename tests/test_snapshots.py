from copy import copy
from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, Mock, patch

from openpyxl import load_workbook
from test_desktop_downloads import report, RELEASE, asset


class SnapshotTests(unittest.TestCase):
    def test_publishing_uses_downloaded_history_and_adds_one_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            existing = root / 'existing.xlsx'
            report.write_report(existing, RELEASE, [asset('DesktopEditors_x64.exe', 10)],
                                datetime(2026, 10, 20, tzinfo=timezone.utc))
            portal = Mock()
            portal.find.return_value = {'id': 5, 'version': 1}
            portal.download.side_effect = lambda path, info: path.write_bytes(existing.read_bytes())
            latest = (RELEASE, [asset('DesktopEditors_x64.exe', 100)])
            previous = (dict(RELEASE, tag_name='v9.4.0'), [asset('DesktopEditors_x64.exe', 20)])
            output = root / 'ONLYOFFICE-Download-Statistics.xlsx'
            with patch('sys.argv', ['report', '--publish', '--output', str(output)]), \
                    patch.dict('os.environ', {'DOCSPACE_API_KEY': 'test'}), \
                    patch.object(report, 'DocSpace', return_value=portal), \
                    patch.object(report, 'collect_recent', side_effect=[[latest, previous], [latest, previous]]):
                self.assertEqual(report.main(), 0)
            saved = load_workbook(output)
            self.assertEqual(len(saved.sheetnames), 2)
            self.assertEqual(saved.worksheets[0].title, '2026-10-20')
            self.assertEqual(saved.worksheets[0]['E8'].value, 10)
            self.assertEqual(saved.active['E8'].value, 100)
            saved.close()
            portal.download.assert_called_once()
            portal.publish.assert_called_once_with(output, portal.find.return_value)

    def test_failed_history_download_does_not_publish(self):
        portal = Mock()
        portal.find.return_value = {'id': 5, 'version': 1}
        portal.download.side_effect = report.requests.ConnectionError('Download failed')
        with patch('sys.argv', ['report', '--publish']), \
                patch.dict('os.environ', {'DOCSPACE_API_KEY': 'test'}), \
                patch.object(report, 'DocSpace', return_value=portal), \
                patch.object(report, 'collect_recent') as collect, self.assertLogs(report.LOG, level='ERROR'):
            self.assertEqual(report.main(), 1)
        collect.assert_not_called()
        portal.publish.assert_not_called()

    def test_selects_previous_stable_release_by_publication_date(self):
        previous = dict(RELEASE, id=9, tag_name='v9.3.1', published_at='2026-09-01T00:00:00Z')
        older = dict(previous, id=8, tag_name='v9.2.0', published_at='2026-08-01T00:00:00Z')
        beta = dict(previous, id=11, tag_name='v10.1.0-beta', prerelease=True,
                    published_at='2026-10-10T00:00:00Z')
        draft = dict(previous, id=12, draft=True)
        latest = (RELEASE, [asset('DesktopEditors_x64.exe', 100)])
        old = (previous, [asset('DesktopEditors_x64.exe', 40)])
        with patch.object(report, 'collect', side_effect=[latest, old]) as collect, \
                patch.object(report, 'request_json', return_value=[older, draft, beta, previous, RELEASE]):
            self.assertEqual(report.collect_recent(Mock()), [latest, old])
        self.assertEqual(collect.call_args.args[1], 'v9.3.1')

    def test_single_stable_release_does_not_invent_a_previous_version(self):
        latest = (RELEASE, [asset('DesktopEditors_x64.exe', 100)])
        with patch.object(report, 'collect', return_value=latest), \
                patch.object(report, 'request_json', return_value=[RELEASE]):
            self.assertEqual(report.collect_recent(Mock()), [latest])

    def test_release_listing_is_paginated(self):
        previous = dict(RELEASE, id=9, tag_name='v9.3.1', published_at='2026-09-01T00:00:00Z')
        latest = (RELEASE, [asset('DesktopEditors_x64.exe', 100)])
        with patch.object(report, 'collect', side_effect=[latest, (previous, latest[1])]), \
                patch.object(report, 'request_json', side_effect=[[dict(RELEASE, draft=True)] * 100,
                                                                   [previous]]) as request:
            self.assertEqual(len(report.collect_recent(Mock())), 2)
        self.assertEqual(request.call_args.kwargs['params']['page'], 2)

    def test_four_versions_are_separate_on_one_sheet(self):
        now = datetime(2026, 10, 28, 9, 17, tzinfo=timezone.utc)
        desktop_previous = (dict(RELEASE, tag_name='v9.3.1'), [asset('DesktopEditors_x64.exe', 20)])
        server = (dict(RELEASE, tag_name='v10.0.1'), [asset('onlyoffice-documentserver.exe', 30)])
        server_previous = (dict(RELEASE, tag_name='v9.3.0'), [asset('onlyoffice-documentserver.exe', 40)])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            report.write_report(path, RELEASE, [asset('DesktopEditors_x64.exe', 10)], now,
                                server=server, previous=desktop_previous, server_previous=server_previous)
            with_values = load_workbook(path, data_only=True)
            sheet = with_values.active
            self.assertEqual(with_values.sheetnames, ['2026-10-28'])
            left = [cell.value for row in sheet for cell in row if cell.column == 1 and cell.value]
            right = [cell.value for row in sheet for cell in row if cell.column == 7 and cell.value]
            self.assertIn('Desktop Editors v10.0.0 — Latest', left)
            self.assertIn('Desktop Editors v9.3.1 — Previous', left)
            self.assertIn('Document Server v10.0.1 — Latest', right)
            self.assertIn('Document Server v9.3.0 — Previous', right)
            self.assertEqual([sheet.cell(row, 5).value for row in range(1, sheet.max_row + 1)
                              if sheet.cell(row, 1).value == 'Total — installations'], [10, 20])
            self.assertEqual([sheet.cell(row, 11).value for row in range(1, sheet.max_row + 1)
                              if sheet.cell(row, 7).value == 'Total — installations'], [30, 40])
            with_values.close()

    def test_appending_preserves_old_values_formulas_and_styles(self):
        now = datetime(2026, 10, 28, 9, 17, tzinfo=timezone.utc)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            report.write_report(path, RELEASE, [asset('DesktopEditors_x64.exe', 10)], now)
            original = load_workbook(path)
            old = {(cell.coordinate): (cell.value, copy(cell.font), copy(cell.fill), cell.number_format)
                   for row in original.active for cell in row}
            original.close()
            report.write_report(path, RELEASE, [asset('DesktopEditors_x64.exe', 99)], now, history_path=path)
            report.write_report(path, RELEASE, [asset('DesktopEditors_x64.exe', 101)], now, history_path=path)
            updated = load_workbook(path)
            self.assertEqual(updated.sheetnames, ['2026-10-28', '2026-10-28 09-17-00', '2026-10-28 09-17-00-2'])
            for coordinate, expected in old.items():
                cell = updated.worksheets[0][coordinate]
                self.assertEqual((cell.value, copy(cell.font), copy(cell.fill), cell.number_format), expected)
            self.assertEqual(updated.active.title, updated.sheetnames[-1])
            self.assertEqual(updated.active['E8'].value, 101)
            self.assertEqual(updated.calculation.calcMode, 'auto')
            updated.close()

    def test_corrupt_history_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            path.write_bytes(b'not a valid workbook')
            with self.assertRaisesRegex(report.ReportError, 'history will not be overwritten'):
                report.write_report(path, RELEASE, [asset('DesktopEditors_x64.exe', 1)],
                                    datetime.now(timezone.utc), history_path=path)
            self.assertEqual(path.read_bytes(), b'not a valid workbook')

    def test_signed_download_does_not_forward_the_api_key(self):
        portal = report.DocSpace('https://nct2.onlyoffice.com', 'secret', '42')
        portal.call = Mock(return_value='https://storage.example.com/signed-report')
        response = MagicMock()
        response.__enter__.return_value.iter_content.return_value = [b'workbook bytes']
        with tempfile.TemporaryDirectory() as directory, patch.object(report.requests, 'get', return_value=response) as get:
            path = Path(directory) / 'history.xlsx'
            portal.download(path, {'id': 5})
            self.assertEqual(path.read_bytes(), b'workbook bytes')
        self.assertNotIn('headers', get.call_args.kwargs)
        self.assertNotIn('auth', get.call_args.kwargs)


if __name__ == '__main__':
    unittest.main()
